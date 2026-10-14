from pathlib import Path
import re, json, csv, html, base64, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '交付'
SOURCE = Path('/Users/zhangmingqi/Desktop/ai情景化训练')
MD = OUT / '开口迎客_独立设计规范_V1.0.md'
text = MD.read_text()
chapters = [x.strip() for x in text.split('<!-- PAGE -->')]

def blocks(source):
    lines=source.splitlines(); i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s: i+=1; continue
        if s.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',x) for x in row): rows.append(row)
                i+=1
            yield 'table', rows
        elif s.startswith('#'):
            n=len(s)-len(s.lstrip('#')); yield f'h{n}',s[n:].strip(); i+=1
        else:
            para=[s]; i+=1
            while i<len(lines) and lines[i].strip() and not lines[i].strip().startswith(('#','|')):
                para.append(lines[i].strip()); i+=1
            yield 'p',' '.join(para)

def runfont(run,size=None,bold=None):
    run.font.name='Arial'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    run.font.color.rgb=RGBColor(0,0,0)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.6); sec.bottom_margin=Cm(1.6)
sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
sec.header_distance=Cm(.55); sec.footer_distance=Cm(.65)
for name,sz in [('Normal',10.5),('Title',25),('Heading 1',19),('Heading 2',12.5)]:
    st=doc.styles[name]; st.font.name='Arial'; st.font.size=Pt(sz); st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    st.paragraph_format.space_after=Pt(6 if name=='Normal' else 9)
    st.paragraph_format.line_spacing=1.35
    if name!='Normal': st.paragraph_format.keep_with_next=True
for b in doc.styles.element.xpath('.//w:pBdr'):
    b.getparent().remove(b)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('开口迎客设计规范 V1.0   '); runfont(r,8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)

for ci,chapter in enumerate(chapters):
    if ci: doc.add_page_break()
    for kind,value in blocks(chapter):
        if kind in ('h1','h2'):
            p=doc.add_paragraph(value, 'Title' if ci==0 and kind=='h1' else 'Heading 1' if kind=='h1' else 'Heading 2')
            for r in p.runs: runfont(r)
        elif kind=='p':
            p=doc.add_paragraph(value)
            p.paragraph_format.widow_control=True
            for r in p.runs: runfont(r)
        elif kind=='table':
            n=len(value[0]); table=doc.add_table(rows=0, cols=n)
            table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
            widths={2:[4.1,13.5],3:[3.4,7.5,6.7],4:[3.1,5.4,5.7,3.4]}.get(n,[17.6/n]*n)
            if ci == 1: widths=[6.6,3.0,8.0]
            if ci==1: widths=[6.6,3.0,8.0]
            for c,w in zip(table.columns,widths): c.width=Cm(w)
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                x=OxmlElement(f'w:{edge}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
            table._tbl.tblPr.append(borders)
            for ri,row in enumerate(value):
                cells=table.add_row().cells
                for j,(c,content) in enumerate(zip(cells,row)):
                    c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcPr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
                    for edge in ('top','left','bottom','right'):
                        m=OxmlElement(f'w:{edge}'); m.set(qn('w:w'),'85'); m.set(qn('w:type'),'dxa'); margins.append(m)
                    tcPr.append(margins)
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6EEF1' if ri==0 else 'F6F8F9' if ri%2==0 else 'FFFFFF'); tcPr.append(shade)
                    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1); p.paragraph_format.line_spacing=1.2
                    if j==0 and (ci in (4,5) or len(content)<10): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    r=p.add_run(content); runfont(r,9.4,ri==0)
                trPr=cells[0]._tc.getparent().get_or_add_trPr()
                cant=OxmlElement('w:cantSplit'); trPr.append(cant)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader'); trPr.append(repeat)
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1
            p.add_run().font.size=Pt(2)
doc.core_properties.title='开口迎客设计规范'
doc.core_properties.subject='产品结构 页面交互 系统架构 设计规范'
doc.core_properties.author='开口迎客项目组'
doc.save(OUT/'开口迎客_独立设计规范_V1.0.docx')
