from pathlib import Path
import re, json, csv, html, base64, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '交付'
SOURCE = Path('/Users/zhangmingqi/Desktop/ai情景化训练')
MD = OUT / '开口迎客_系统架构与设计规范_V3.0.md'
text = MD.read_text()
chapters = [x.strip() for x in text.split('<!-- PAGE -->')]

def blocks(source):
    lines=source.splitlines(); i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s: i+=1; continue
        if s.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',x) for x in row): rows.append(row)
                i+=1
            yield 'table', rows
        elif s.startswith('#'):
            n=len(s)-len(s.lstrip('#')); yield f'h{n}',s[n:].strip(); i+=1
        else:
            para=[s]; i+=1
            while i<len(lines) and lines[i].strip() and not lines[i].strip().startswith(('#','|')):
                para.append(lines[i].strip()); i+=1
            yield 'p',' '.join(para)

def runfont(run,size=None,bold=None):
    run.font.name='Arial'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    run.font.color.rgb=RGBColor(0,0,0)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.6); sec.bottom_margin=Cm(1.6)
sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
sec.header_distance=Cm(.55); sec.footer_distance=Cm(.65)
for name,sz in [('Normal',10.5),('Title',25),('Heading 1',19),('Heading 2',12.5)]:
    st=doc.styles[name]; st.font.name='Arial'; st.font.size=Pt(sz); st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    st.paragraph_format.space_after=Pt(6 if name=='Normal' else 9)
    st.paragraph_format.line_spacing=1.35
    if name!='Normal': st.paragraph_format.keep_with_next=True
for b in doc.styles.element.xpath('.//w:pBdr'):
    b.getparent().remove(b)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('开口迎客 V3.0   '); runfont(r,8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)

for ci,chapter in enumerate(chapters):
    if ci: doc.add_page_break()
    for kind,value in blocks(chapter):
        if kind in ('h1','h2'):
            p=doc.add_paragraph(value, 'Title' if ci==0 and kind=='h1' else 'Heading 1' if kind=='h1' else 'Heading 2')
            for r in p.runs: runfont(r)
        elif kind=='p':
            p=doc.add_paragraph(value)
            p.paragraph_format.widow_control=True
            for r in p.runs: runfont(r)
        elif kind=='table':
            n=len(value[0]); table=doc.add_table(rows=0, cols=n)
            table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
            widths={2:[4.1,13.5],3:[3.4,7.5,6.7],4:[3.1,5.4,5.7,3.4]}.get(n,[17.6/n]*n)
            if ci in (4,5): widths=[1.65,7.9,8.05]
            if ci==16: widths=[6.6,3.0,8.0]
            for c,w in zip(table.columns,widths): c.width=Cm(w)
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                x=OxmlElement(f'w:{edge}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
            table._tbl.tblPr.append(borders)
            for ri,row in enumerate(value):
                cells=table.add_row().cells
                for j,(c,content) in enumerate(zip(cells,row)):
                    c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcPr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
                    for edge in ('top','left','bottom','right'):
                        m=OxmlElement(f'w:{edge}'); m.set(qn('w:w'),'65' if ci in (4,5) else '85'); m.set(qn('w:type'),'dxa'); margins.append(m)
                    tcPr.append(margins)
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6EEF1' if ri==0 else 'F6F8F9' if ri%2==0 else 'FFFFFF'); tcPr.append(shade)
                    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1); p.paragraph_format.line_spacing=1.2
                    if j==0 and (ci in (4,5) or len(content)<10): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    r=p.add_run(content); runfont(r,8.8 if ci in (4,5) else 9.4,ri==0)
                trPr=cells[0]._tc.getparent().get_or_add_trPr()
                cant=OxmlElement('w:cantSplit'); trPr.append(cant)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader'); trPr.append(repeat)
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1
            p.add_run().font.size=Pt(2)
doc.core_properties.title='开口迎客系统架构与设计规范'
doc.core_properties.subject='产品结构 页面交互 系统架构 设计规范'
doc.core_properties.author='开口迎客项目组'
doc.save(OUT/'开口迎客_系统架构与设计规范_V3.0.docx')

routes=[]
for line in text.splitlines():
    m=re.match(r'\| ([A-Z]\d{2}) ([PE]) \| (.*?) (/[\w/:.-]+) \| (.*?) \|',line)
    if m: routes.append(dict(zip(['id','phase','name','route','primary_action'],m.groups())))
assert len(routes)==47, len(routes)
with (OUT/'页面路由与交付清单.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=routes[0].keys()); w.writeheader(); w.writerows(routes)

colors={}
for line in text.splitlines():
    m=re.match(r'\| (color\.[a-z.]+) \| (#[0-9A-F]{6}) \| (.*?) \|',line)
    if m: colors[m[1]]={'value':m[2],'description':m[3]}
tokens={'version':'3.0','status':'proposed','color':colors,
        'typography':{'unit':'sp','title':28,'pageTitle':24,'cardTitle':22,'body':20,'supporting':18,'tabLabel':14,'lineHeight':1.5},
        'spacing':{'unit':'dp','scale':[4,8,12,16,24,32]},
        'radius':{'card':20,'button':28},
        'layout':{'unit':'dp','baseWidth':390,'testWidths':[360,390,412],'pagePadding':16,'cardPadding':[16,20],'primaryButtonHeight':[56,64],'minimumTouchTarget':48,'recorderTarget':72},
        'motion':{'unit':'ms','transitionRange':[150,250],'respectsReducedMotion':True}}
(OUT/'设计Tokens.json').write_text(json.dumps(tokens,ensure_ascii=False,indent=2))

def luminance(h):
    c=[int(h[i:i+2],16)/255 for i in (1,3,5)]
    c=[v/12.92 if v<=.04045 else ((v+.055)/1.055)**2.4 for v in c]
    return .2126*c[0]+.7152*c[1]+.0722*c[2]
contrast={k:round(1.05/(luminance(v['value'])+.05),3) for k,v in colors.items()}

# HTML is an offline design review document, not an app implementation.
def esc(x): return html.escape(str(x),quote=True)
def inline(x):
    x=esc(x)
    return re.sub(r'(https://[^\s<]+)',r'<a href="\1" target="_blank" rel="noopener">\1</a>',x)
def mdhtml(c):
    out=[]
    for k,v in blocks(c):
        if k=='table': out.append('<div class="table-wrap"><table><thead><tr>'+''.join('<th>'+inline(x)+'</th>' for x in v[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+inline(x)+'</td>' for x in row)+'</tr>' for row in v[1:])+'</tbody></table></div>')
        else: out.append(f'<{k}>{inline(v)}</{k}>')
    return ''.join(out)

refs=[('首页','微信图片_20261014094946.png'),('情境课程','微信图片_20261014095004.png'),('AI游客','微信图片_20261014095000.png'),('复习','微信图片_20261014094955.png'),('我的','微信图片_20261014094950.png')]
def data_image(name): return 'data:image/png;base64,'+base64.b64encode((SOURCE/name).read_bytes()).decode()
gallery=''.join(f'<figure><img loading="lazy" src="{data_image(f)}" alt="原始{n}页面参考"><figcaption>{n}<small>原始视觉基线</small></figcaption></figure>' for n,f in refs)
swatches=''.join(f'<div class="swatch"><div style="background:{v["value"]}"></div><b>{v["value"]}</b><span>{esc(v["description"])}</span><code>{k}</code></div>' for k,v in colors.items())
nav=''.join(f'<a href="#chapter-{i}">{esc(c.splitlines()[0].lstrip("# "))}</a>' for i,c in enumerate(chapters) if i)
manual=''.join(f'<article class="chapter" id="chapter-{i}">{mdhtml(c)}</article>' for i,c in enumerate(chapters))
routehtml=''.join(f'<tr><td><code>{x["id"]}</code></td><td>{x["name"]}</td><td>{"首期" if x["phase"]=="P" else "扩展"}</td><td><code>{x["route"]}</code></td><td>{x["primary_action"]}</td></tr>' for x in routes)
base=(ROOT/'scripts/design_overview.html').read_text()
for key,value in {'GALLERY':gallery,'SWATCHES':swatches,'CHAPTER_NAV':nav,'MANUAL':manual,'ROUTES':routehtml,'CONTRAST':json.dumps(contrast,ensure_ascii=False)}.items(): base=base.replace('{{'+key+'}}',value)
(OUT/'设计总览.html').write_text(base)
(OUT/'README.md').write_text('''# 开口迎客 V3.0 交付说明

先打开 `设计总览.html` 查看系统结构、原始页面、补充页面示意、组件与完整规范。该文件可离线打开，无外部脚本和网络请求。补充页面为静态设计示意，不提供真实录音、AI或后台功能。

- `开口迎客_系统架构与设计规范_V3.0.docx` 为可编辑评审手册。
- 同名 `.md` 为内容源文件，方便持续维护。
- `页面路由与交付清单.csv` 包含47个页面定义，46个首期、1个扩展。
- `设计Tokens.json` 为建议的视觉与布局变量。
- `系统架构与学习闭环.md` 为包含5张Mermaid图的可编辑图源。

保留5张参考图的视觉语言与导航，已确认“我想说”固定在首页。本版是完整架构与设计方案，未实现生产App或连接模型。MVP范围收敛、教学阈值、数据保留参数与供应商能力待项目组评审。原始文件未修改。

生成入口为 `scripts/build_deliverables.py`，运行时需使用Codex工作区附带的Python依赖。视觉总览及Word均由同一内容源生成。具体设备、地区语言与经营事实在真实试点中确认。
''')
manifest={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in SOURCE.iterdir() if f.is_file()}
(ROOT/'.work').mkdir(exist_ok=True)
(ROOT/'.work/source_sha256.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps({'chapters':len(chapters),'pages_defined':len(routes),'tokens':len(colors),'contrast_white':contrast},ensure_ascii=False))
